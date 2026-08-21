#!/usr/bin/env python3
"""Dump a spreadsheet's structure: sheets, cell values, and formulas (xlsx only).

Usage: dump_sheet.py FILE [sheet_index_or_name] [max_rows] [max_cols]
With no sheet arg: lists sheets with dimensions. With a sheet arg: dumps a grid
of non-empty cells as `A1: value | =FORMULA` lines, row by row.
"""
import sys, os

def dump_xlsx(path, sheet=None, max_rows=400, max_cols=40):
    import openpyxl
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    if sheet is None:
        for ws in wbf.worksheets:
            print(f"SHEET {ws.title!r}: {ws.max_row} rows x {ws.max_column} cols")
        return
    wsf = wbf[sheet] if sheet in wbf.sheetnames else wbf.worksheets[int(sheet)]
    wsv = wbv[wsf.title]
    for row in wsf.iter_rows(min_row=1, max_row=min(wsf.max_row, max_rows),
                             max_col=min(wsf.max_column, max_cols)):
        for c in row:
            if c.value is None:
                continue
            v = wsv[c.coordinate].value
            if isinstance(c.value, str) and c.value.startswith("="):
                print(f"{c.coordinate}: {v!r} | {c.value}")
            else:
                print(f"{c.coordinate}: {c.value!r}")

def dump_xls(path, sheet=None, max_rows=400, max_cols=40):
    import xlrd
    wb = xlrd.open_workbook(path, formatting_info=False)
    if sheet is None:
        for ws in wb.sheets():
            print(f"SHEET {ws.name!r}: {ws.nrows} rows x {ws.ncols} cols")
        return
    names = wb.sheet_names()
    ws = wb.sheet_by_name(sheet) if sheet in names else wb.sheet_by_index(int(sheet))
    from xlrd.xldate import xldate_as_datetime
    for r in range(min(ws.nrows, max_rows)):
        for c in range(min(ws.ncols, max_cols)):
            cell = ws.cell(r, c)
            if cell.ctype == 0:
                continue
            coord = f"{chr(65 + c % 26) if c < 26 else chr(64 + c // 26) + chr(65 + c % 26)}{r + 1}"
            print(f"{coord}: {cell.value!r}")

if __name__ == "__main__":
    path = sys.argv[1]
    sheet = sys.argv[2] if len(sys.argv) > 2 else None
    max_rows = int(sys.argv[3]) if len(sys.argv) > 3 else 400
    max_cols = int(sys.argv[4]) if len(sys.argv) > 4 else 40
    if path.lower().endswith(".xlsx"):
        dump_xlsx(path, sheet, max_rows, max_cols)
    else:
        dump_xls(path, sheet, max_rows, max_cols)
